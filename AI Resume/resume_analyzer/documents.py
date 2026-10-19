from __future__ import annotations

from typing import BinaryIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader
from pypdf.errors import PdfReadError, PdfStreamError


class DocumentError(ValueError):
    """Raised when an uploaded document cannot be safely read or has no text."""


def extract_text(file_obj: BinaryIO, extension: str) -> str:
    """Extract readable text from a supported PDF, DOCX, or plain-text file."""
    suffix = extension.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(file_obj)
            if reader.is_encrypted:
                raise DocumentError("Password-protected PDF files are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            document = Document(file_obj)
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                parts.extend(
                    cell.text for row in table.rows for cell in row.cells
                )
            text = "\n".join(parts)
        elif suffix == ".txt":
            text = file_obj.read().decode("utf-8-sig")
        else:
            raise DocumentError(
                f"Unsupported file type {suffix!r}. "
                "Supported types are PDF, DOCX, and TXT."
            )
    except DocumentError:
        raise
    except (
        BadZipFile,
        UnicodeDecodeError,
        OSError,
        ValueError,
        KeyError,
        PackageNotFoundError,
        PdfReadError,
        PdfStreamError,
        XMLSyntaxError,
    ) as error:
        raise DocumentError(
            "The document could not be read. Check that it is a valid, "
            "unprotected PDF, DOCX, or UTF-8 text file."
        ) from error

    if not text.strip():
        raise DocumentError(
            "No readable text was found. Scanned/image-only PDFs need OCR "
            "before they can be analyzed."
        )
    return text
